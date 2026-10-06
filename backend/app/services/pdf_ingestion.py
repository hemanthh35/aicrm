"""Bounded extraction and chunking for administrator-uploaded text PDFs."""

from io import BytesIO
import re
from zipfile import ZipFile

from docx import Document
from pypdf import PdfReader

MAX_PDF_PAGES = 100
MAX_EXTRACTED_CHARS = 300_000
MAX_CHUNK_CHARS = 2_400
CHUNK_OVERLAP = 300
MAX_DOCX_MEMBERS = 1_000


class PdfExtractionError(ValueError):
    """Raised when a PDF cannot be safely converted into searchable text."""


def _chunk_page(text: str) -> list[str]:
    normalized = re.sub(r"[ \t]+", " ", text.replace("\x00", " ")).strip()
    if not normalized:
        return []

    chunks: list[str] = []
    start = 0
    while start < len(normalized):
        hard_end = min(start + MAX_CHUNK_CHARS, len(normalized))
        end = hard_end
        if hard_end < len(normalized):
            boundaries = [
                normalized.rfind("\n\n", start + MAX_CHUNK_CHARS // 2, hard_end),
                normalized.rfind(". ", start + MAX_CHUNK_CHARS // 2, hard_end),
                normalized.rfind(" ", start + MAX_CHUNK_CHARS // 2, hard_end),
            ]
            boundary = max(boundaries)
            if boundary > start:
                end = boundary + (2 if normalized[boundary:boundary + 2] == ". " else 0)

        chunk = normalized[start:end].strip()
        if chunk:
            chunks.append(chunk)
        if end >= len(normalized):
            break
        start = max(end - CHUNK_OVERLAP, start + 1)
    return chunks


def extract_pdf_chunks(contents: bytes) -> tuple[int, list[tuple[int, str]]]:
    if not contents.startswith(b"%PDF-"):
        raise PdfExtractionError("The uploaded file is not a valid PDF")

    try:
        reader = PdfReader(BytesIO(contents), strict=False)
        page_count = len(reader.pages)
    except Exception as exc:  # pypdf raises several parser-specific exceptions
        raise PdfExtractionError("The PDF could not be read") from exc

    if page_count < 1:
        raise PdfExtractionError("The PDF has no pages")
    if page_count > MAX_PDF_PAGES:
        raise PdfExtractionError(f"PDFs are limited to {MAX_PDF_PAGES} pages")

    chunks: list[tuple[int, str]] = []
    extracted_chars = 0
    try:
        for page_number, page in enumerate(reader.pages, start=1):
            page_text = page.extract_text() or ""
            extracted_chars += len(page_text)
            if extracted_chars > MAX_EXTRACTED_CHARS:
                raise PdfExtractionError("The PDF contains too much text to index")
            chunks.extend((page_number, chunk) for chunk in _chunk_page(page_text))
    except PdfExtractionError:
        raise
    except Exception as exc:
        raise PdfExtractionError("The PDF text could not be extracted") from exc

    if not chunks:
        raise PdfExtractionError("This PDF has no extractable text. Scanned PDFs need OCR before upload")
    return page_count, chunks


def extract_docx_chunks(contents: bytes) -> tuple[int, list[tuple[int, str]]]:
    """Extract bounded paragraph/table text from a DOCX without unpacking it to disk."""
    if not contents.startswith(b"PK"):
        raise PdfExtractionError("The uploaded file is not a valid DOCX")
    try:
        with ZipFile(BytesIO(contents)) as archive:
            members = archive.infolist()
            if len(members) > MAX_DOCX_MEMBERS:
                raise PdfExtractionError("The DOCX contains too many archive entries")
            total_uncompressed = sum(member.file_size for member in members)
            if total_uncompressed > MAX_EXTRACTED_CHARS * 4:
                raise PdfExtractionError("The DOCX contains too much uncompressed content")
            if any(member.filename.startswith(("/", "\\")) or ".." in member.filename.split("/") for member in members):
                raise PdfExtractionError("The DOCX contains unsafe archive paths")
        document = Document(BytesIO(contents))
    except PdfExtractionError:
        raise
    except Exception as exc:
        raise PdfExtractionError("The DOCX could not be read") from exc

    parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                parts.append(row_text)
    chunks = _chunk_page("\n\n".join(parts))
    if not chunks:
        raise PdfExtractionError("This DOCX has no extractable text")
    return 1, [(1, chunk) for chunk in chunks]
