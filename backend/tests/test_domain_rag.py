from io import BytesIO

import pytest
from docx import Document
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from app.db.session import Base
from app.models.models import Customer, Domain, KnowledgeBase, KnowledgeChunk, KnowledgeDocument, Language
from app.services import rag, vector_index
from app.services.ai_guardrails import normalize_assistant_result
from app.services.pdf_ingestion import PdfExtractionError, extract_docx_chunks, extract_pdf_chunks
from app.routers.assistant import AGRIMART_SYSTEM_PROMPT, TECHKART_SYSTEM_PROMPT


def _make_pdf(text: bytes) -> bytes:
    content = b"BT\n/F1 12 Tf\n72 720 Td\n(" + text + b") Tj\nET\n"
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>",
        b"<< /Length " + str(len(content)).encode() + b" >>\nstream\n" + content + b"endstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    pdf = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for index, obj in enumerate(objects, 1):
        offsets.append(len(pdf))
        pdf.extend(f"{index} 0 obj\n".encode())
        pdf.extend(obj)
        pdf.extend(b"\nendobj\n")
    xref = len(pdf)
    pdf.extend(f"xref\n0 {len(objects) + 1}\n".encode())
    pdf.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        pdf.extend(f"{offset:010d} 00000 n \n".encode())
    pdf.extend(f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode())
    return bytes(pdf)


def test_domain_prompts_are_distinct_and_agriculture_specific():
    assert "TechKart" in TECHKART_SYSTEM_PROMPT
    assert "AgriMart" in AGRIMART_SYSTEM_PROMPT
    assert "pest_or_disease" in AGRIMART_SYSTEM_PROMPT
    assert "fertilizer doses" in AGRIMART_SYSTEM_PROMPT
    assert TECHKART_SYSTEM_PROMPT != AGRIMART_SYSTEM_PROMPT


def test_agriculture_intent_schema_is_normalized_separately():
    result = normalize_assistant_result({"reply": "Inspect the leaves.", "intent": "pest_or_disease"}, domain="agriculture")
    assert result["intent"] == "pest_or_disease"

    ecommerce_result = normalize_assistant_result({"reply": "Inspect the leaves.", "intent": "pest_or_disease"})
    assert ecommerce_result["intent"] == "general_query"


def test_pdf_extraction_returns_page_aware_chunks():
    page_count, chunks = extract_pdf_chunks(_make_pdf(b"AgriMart crop guidance and irrigation"))
    assert page_count == 1
    assert chunks[0][0] == 1
    assert "crop guidance" in chunks[0][1]


def test_pdf_extraction_rejects_invalid_or_unsearchable_files():
    with pytest.raises(PdfExtractionError, match="valid PDF"):
        extract_pdf_chunks(b"not a pdf")

    blank_pdf = _make_pdf(b"")
    with pytest.raises(PdfExtractionError, match="extractable text"):
        extract_pdf_chunks(blank_pdf)


def test_docx_extraction_returns_searchable_text():
    document = Document()
    document.add_paragraph("AgriMart soil and irrigation guidance")
    output = BytesIO()
    document.save(output)
    page_count, chunks = extract_docx_chunks(output.getvalue())
    assert page_count == 1
    assert "soil and irrigation" in chunks[0][1]


def test_pdf_chunks_are_retrieved_only_for_their_domain(monkeypatch):
    engine = create_engine("sqlite:///:memory:", connect_args={"check_same_thread": False})
    Base.metadata.create_all(bind=engine)
    session = sessionmaker(bind=engine)()
    try:
        customer = Customer(name="Admin", email="rag-test@example.com", password_hash="")
        session.add(customer)
        session.flush()
        document = KnowledgeDocument(
            original_filename="field-guide.pdf",
            content_type="application/pdf",
            size_bytes=123,
            page_count=1,
            domain=Domain.agriculture,
            language=Language.en,
            created_by=customer.id,
        )
        article = KnowledgeBase(
            title="Warranty guide",
            content="Laptop warranty coverage",
            domain=Domain.ecommerce,
            language=Language.en,
        )
        session.add_all([document, article])
        session.flush()
        session.add(KnowledgeChunk(document_id=document.id, chunk_index=0, page_number=1, content="Crop irrigation guidance"))
        session.commit()

        vector = [1.0] + [0.0] * 383
        monkeypatch.setattr(vector_index, "embed_texts", lambda texts: [vector for _ in texts])
        vector_index.rebuild_index(session)

        agri_hits = rag.retrieve_knowledge(session, "irrigation", Domain.agriculture, Language.en)
        tech_hits = rag.retrieve_knowledge(session, "warranty", Domain.ecommerce, Language.en)
        assert agri_hits and agri_hits[0].source_type == "pdf"
        assert agri_hits[0].page_number == 1
        assert tech_hits and tech_hits[0].source_type == "article"
        assert all(hit.domain == Domain.ecommerce for hit in tech_hits)
    finally:
        vector_index._index = None
        vector_index._id_map = []
        session.close()
